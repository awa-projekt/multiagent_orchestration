"""Word documents (`docx`): outline, tables, comments, tracked changes, hidden text, embedded
parts; nothing in the file is executed, opened or fetched, and its unpacked size is bounded."""

from __future__ import annotations

import html
import re
import zipfile
import zlib
from pathlib import Path
from typing import Any

from awa.data._injection import guess_language, injection_flags
from awa.data._model import DocxInfo, DocxTable, FileProfile, Heading
from awa.data._parse import ParseContext

# What reading a damaged zip member raises besides `BadZipFile`: a broken deflate stream
# (`zlib.error`), a bad offset or length (`OSError`, `EOFError`), an unknown compression
# method (`NotImplementedError`) or an encrypted member (`RuntimeError`).
_DAMAGED_ZIP = (
    zipfile.BadZipFile,
    zlib.error,
    OSError,
    EOFError,
    NotImplementedError,
    RuntimeError,
)

KINDS = ("docx",)

EXCERPT_CHARS = 1500
MAX_HEADINGS = 60
MAX_TABLES = 50
_MAX_HEADER_CELLS = 30
_HEAD_TEXT_CHARS = 20_000

_OLE_MAGIC = b"\xd0\xcf\x11\xe0"
_ENCRYPTED_STREAM = "EncryptedPackage".encode("utf-16-le")
_TEXT_PART = re.compile(r"word/(?:document|comments|footnotes|endnotes|header\d*|footer\d*)\.xml")
_HEADER_FOOTER = re.compile(r"word/(?:header|footer)\d*\.xml")
_RUN_TEXT = re.compile(r"<w:(?:t|delText)\b[^<>]*>([^<]*)</w:(?:t|delText)>|(</w:p>)")
_TRACKED = re.compile(r"<w:(?:ins|del|moveFrom|moveTo)\b")
_COMMENT = re.compile(r"<w:comment\b")
_RELATIONSHIP = re.compile(r"<Relationship\b([^<>]*)>")
_ATTRIBUTE = re.compile(r'([\w:.-]+)\s*=\s*"([^"<>]*)"')
_HEADING = re.compile(r"Heading (\d)")
_HIDDEN_RUNS = (
    ".//w:r[w:t][w:rPr/w:vanish[not(@w:val='0' or @w:val='false')]]",
    ".//w:r[w:t][w:rPr/w:sz[@w:val<=2]]",
    ".//w:r[w:t][w:rPr/w:color[@w:val='FFFFFF']][not(w:rPr/w:shd)]"
    "[not(ancestor::w:tc/w:tcPr/w:shd)][not(ancestor::w:p/w:pPr/w:shd)]",
)


def profile_docx(path: Path, ctx: ParseContext) -> FileProfile:
    """Profile a Word document: outline, tables, text excerpt and the facts that need care."""

    if ctx.library_missing("docx"):
        return ctx.profile
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.oxml.ns import qn

    with path.open("rb") as handle:
        old_format = handle.read(4) == _OLE_MAGIC
    if old_format:
        if _ENCRYPTED_STREAM in path.read_bytes():
            return ctx.fail("DF-E-ENCRYPTED", "the Word file is password protected")
        return ctx.fail("DF-E-PARSE-FAILED", "an old binary Word file (.doc); save it as .docx")
    try:
        names, parts = _read_parts(path, ctx)
    except _DAMAGED_ZIP:
        return ctx.fail("DF-E-PARSE-FAILED", "the file is not a valid Word document (damaged zip)")
    except ValueError as error:
        return ctx.fail("DF-E-ZIP-UNSAFE", f"the Word file {error}")
    try:
        document = Document(str(path))
    except (PackageNotFoundError, KeyError, SyntaxError, *_DAMAGED_ZIP):  # lxml's XMLSyntaxError
        return ctx.fail(
            "DF-E-PARSE-FAILED", "the file is not a valid Word document (or its XML is hostile)"
        )
    info = ctx.profile.docx = DocxInfo(paragraphs=len(document.paragraphs))
    body: Any = document.element.body
    _outline(document, info, ctx)
    info.text_length = sum(len(node.text or "") for node in body.iter(qn("w:t")))
    _archive_facts(names, parts, info, ctx)
    hidden = sum(len(body.xpath(query)) for query in _HIDDEN_RUNS)
    if hidden:
        ctx.flag("DF-W-HIDDEN-TEXT", f"{hidden} text runs are hidden, tiny or white on white")
    text = "\n".join(_part_text(xml) for name, xml in parts.items() if name.endswith(".xml"))
    info.language = guess_language(text[:_HEAD_TEXT_CHARS])
    ctx.profile.flags.extend(injection_flags(text, ctx.file_id))
    return ctx.profile


def _read_parts(path: Path, ctx: ParseContext) -> tuple[list[str], dict[str, str]]:
    """All part names and the text parts; refused when it unpacks beyond the single-file limit."""

    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        total = sum(entry.file_size for entry in entries)
        if total > ctx.limits.max_file_bytes:
            raise ValueError(f"unpacks to {total} bytes; the limit is {ctx.limits.max_file_bytes}")
        parts = {
            entry.filename: archive.read(entry).decode("utf-8", "replace")
            for entry in entries
            if _TEXT_PART.fullmatch(entry.filename) or entry.filename.endswith(".rels")
        }
    return [entry.filename for entry in entries], parts


def _part_text(xml: str) -> str:
    """The text of an XML part, paragraphs separated by newlines (entities resolved)."""

    pieces = ("\n" if paragraph_end else text for text, paragraph_end in _RUN_TEXT.findall(xml))
    return html.unescape("".join(pieces))


def _outline(document: Any, info: DocxInfo, ctx: ParseContext) -> None:
    """Headings, tables and the reading-order excerpt, one pass over the body."""

    limit = ctx.limits.max_chars_per_cell
    lines: list[str] = []
    size = 0
    cut = False
    unreadable = 0
    for count, block in enumerate(document.iter_inner_content()):
        if count % 200 == 0:
            ctx.check_deadline()
        try:
            if hasattr(block, "rows"):
                first = _row_texts(block.rows[0]) if block.rows else []
                header = [cell[:limit] for cell in first]
                table = DocxTable(rows=len(block.rows), columns=len(block.columns), header=header)
                line = " | ".join(first)
                if len(info.tables) < MAX_TABLES:
                    info.tables.append(table)
                else:
                    cut = True
            else:
                line = block.text.strip()
                style = (block.style.name if block.style is not None else "") or ""
                level = _HEADING.fullmatch(style)
                if line and level:
                    if len(info.headings) < MAX_HEADINGS:
                        info.headings.append(Heading(level=int(level[1]), text=line[:limit]))
                    else:
                        cut = True
                    line = f"# {line}"
        except Exception:  # an odd table or style (no grid, broken rows) must not fail the file
            unreadable += 1
            continue
        if line and size < _HEAD_TEXT_CHARS:
            lines.append(line)
            size += len(line)
    info.excerpt = "\n".join(lines)[:EXCERPT_CHARS]
    if unreadable:
        ctx.flag("DF-W-DOCX-ODD-PARTS", f"{unreadable} tables or paragraphs could not be read")
    if cut:
        ctx.profile.status = "partial"
        ctx.flag(
            "DF-W-SAMPLED",
            f"only the first {MAX_HEADINGS} headings and {MAX_TABLES} tables are listed",
        )


def _row_texts(row: Any) -> list[str]:
    cells: list[Any] = []
    for cell in row.cells:  # a merged cell repeats; keep it once
        if not any(cell._tc is seen._tc for seen in cells):
            cells.append(cell)
    return [" ".join(cell.text.split()) for cell in cells[:_MAX_HEADER_CELLS]]


def _archive_facts(
    names: list[str], parts: dict[str, str], info: DocxInfo, ctx: ParseContext
) -> None:
    """Counts and flags that come from the package, not from the text."""

    info.images = sum(name.startswith("word/media/") for name in names)
    info.tracked_changes = len(_TRACKED.findall(parts.get("word/document.xml", "")))
    info.comments = len(_COMMENT.findall(parts.get("word/comments.xml", "")))
    info.headers_footers = any(
        _part_text(xml).strip() for name, xml in parts.items() if _HEADER_FOOTER.fullmatch(name)
    )
    embedded = sum(name.startswith(("word/embeddings/", "word/afchunk")) for name in names)
    if embedded:
        ctx.flag(
            "DF-W-EMBEDDED-OBJECTS", f"{embedded} embedded objects or included documents; not read"
        )
    if any(name.endswith("vbaProject.bin") for name in names):
        ctx.flag("DF-W-MACROS", "the file contains a macro project; it is never run")
    if any(_external(xml) for name, xml in parts.items() if name.endswith(".rels")):
        ctx.flag("DF-W-EXTERNAL-CONTENT", "the file links to outside content; it is never fetched")


def _external(rels: str) -> bool:
    """Whether a relationship part links an external target other than a hyperlink."""

    for tag in _RELATIONSHIP.findall(rels):
        attrs = dict(_ATTRIBUTE.findall(tag))
        if attrs.get("TargetMode") == "External" and not attrs.get("Type", "").endswith(
            "/hyperlink"
        ):
            return True
    return False
