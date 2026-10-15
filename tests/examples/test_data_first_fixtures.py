"""The generated files of the data-first scenarios: reproducible, small, and exactly what the
scenarios say they are (`tests/_helpers/data_scenarios/generate.py`).

The files must be the same bytes on every run of one library set; a scenario test that relied on
a file that silently changed shape would prove nothing. Each content test states the property a
profiler is expected to find, taken from the file itself and not from the profiler.
"""

from __future__ import annotations

import csv
import email
import email.policy
import io
import json
import os
import sqlite3
import struct
import subprocess
import sys
from pathlib import Path
from typing import Any

import pypdf
import pytest

from _helpers.data_scenarios import generate

SUPPORT = Path(__file__).resolve().parents[1] / "_helpers"

GENERATE_PY = SUPPORT / "data_scenarios" / "generate.py"


@pytest.fixture(scope="module")
def out(tmp_path_factory: pytest.TempPathFactory) -> Path:
    root = tmp_path_factory.mktemp("scenario_files")
    generate.generate(root)
    return root


# ------------------------------------------------------------------------ the generator itself


def test_two_processes_with_different_hash_seeds_write_identical_bytes(tmp_path: Path) -> None:
    folders = []
    for seed in ("1", "4242"):
        folder = tmp_path / f"seed{seed}"
        env = {**os.environ, "PYTHONHASHSEED": seed}
        subprocess.run(
            [sys.executable, str(GENERATE_PY), str(folder)],
            env=env,
            check=True,
            capture_output=True,
        )
        folders.append(folder)
    names = sorted(
        p.relative_to(folders[0]).as_posix() for p in folders[0].rglob("*") if p.is_file()
    )
    assert names == sorted(generate.FILES)
    for name in names:
        assert (folders[0] / name).read_bytes() == (folders[1] / name).read_bytes(), name


def test_check_mode_passes_on_a_fresh_folder_and_names_changed_and_missing_files(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    generate.generate(tmp_path)
    assert generate.main(["--check", str(tmp_path)]) == 0
    assert "up to date" in capsys.readouterr().out
    (tmp_path / "csv" / "orders_messy.csv").write_bytes(b"changed")
    (tmp_path / "sqlite" / "shop.db").unlink()
    assert generate.main(["--check", str(tmp_path)]) == 1
    assert capsys.readouterr().out.split() == ["csv/orders_messy.csv", "sqlite/shop.db"]
    assert generate.differing(tmp_path) == ["csv/orders_messy.csv", "sqlite/shop.db"]


def test_the_whole_set_is_small(out: Path) -> None:
    sizes = {p.name: p.stat().st_size for p in out.rglob("*") if p.is_file()}
    assert len(sizes) == len(generate.FILES) == 18
    assert sum(sizes.values()) < 1_000_000


# ------------------------------------------------------------------------ messy German CSV


def test_orders_messy_has_the_defects_the_scenario_names(out: Path) -> None:
    raw = (out / "csv" / "orders_messy.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8-sig").splitlines()
    assert lines[-1] == "Summe;;;123.456,75;;"
    rows = [line.split(";") for line in lines[1:-1] if line]
    assert [len(r) for r in rows].count(3) == 1  # the ragged row
    assert lines.count("") == 2  # blank lines inside the data
    full = [r for r in rows if len(r) > 3]
    assert "4250.00" in {r[3] for r in full}
    # the one amount written with a decimal point (4250.00) is read as it is, the others as German
    amounts = [
        float(r[3]) if "," not in r[3] else float(r[3].replace(".", "").replace(",", "."))
        for r in full
    ]
    assert round(sum(amounts) * 100) == generate.MESSY_TOTAL_CENTS  # the Summe line is the true sum
    dates = {r[4] for r in full}
    assert "04.05.2024" in dates and max(int(d[:2]) for d in dates) > 12  # a day proves day-first


def test_bank_statement_is_cp1252_with_title_lines_and_a_quoted_line_break(out: Path) -> None:
    text = (out / "csv" / "kontoauszug_sparkasse.csv").read_bytes().decode("cp1252")
    rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=";"))
    assert rows[0][0] == "Umsatzanzeige" and rows[2] == []
    header = rows[3]
    assert header[:2] == ["Buchungstag", "Wertstellung"] and header[-1] == "Währung"
    assert any("\n" in cell for row in rows for cell in row)
    assert any(cell == "RE 2024-0412; Müller & Söhne KG" for row in rows for cell in row)
    assert {len(r) for r in rows[4:]} == {6}


def test_the_utf16_file_has_a_bom_tabs_leading_zeros_and_parenthesised_negatives(out: Path) -> None:
    raw = (out / "csv" / "kunden_utf16.csv").read_bytes()
    assert raw[:2] == b"\xff\xfe"
    rows = [line.split("\t") for line in raw.decode("utf-16").splitlines()]
    assert rows[0] == ["Kundennr", "Name", "Rabatt", "Saldo", "Angelegt"]
    assert rows[1][0] == "00123" and rows[1][2] == "12,5 %" and rows[1][3] == "(1.234,56)"


def test_ambiguous_dates_cannot_be_read_either_way(out: Path) -> None:
    lines = (out / "csv" / "termine_mehrdeutig.csv").read_text().splitlines()[1:]
    dates = [line.split(";")[1] for line in lines]
    assert dates and all(int(d[:2]) <= 12 and int(d[3:5]) <= 12 for d in dates)


# ------------------------------------------------------------------------ Excel


def test_the_sales_workbook_has_the_layout_scenario_4_describes(out: Path) -> None:
    import openpyxl

    path = out / "excel" / "s4_sales.xlsx"
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Umsatz 2026", "Ziele", "Notizen"]
    assert [s.sheet_state for s in book] == ["visible", "visible", "hidden"]
    sheet = book["Umsatz 2026"]
    assert [str(r) for r in sheet.merged_cells.ranges] == ["A1:F1"]
    header = [c.value for c in sheet[generate.SALES_HEADER_ROW]]
    assert header == ["Kunde", "Monat", "Bestellungen", "Umsatz_EUR", "Region", "Betreuer"]
    data_rows = list(sheet.iter_rows(min_row=4, max_row=27, values_only=True))
    assert len(data_rows) == 24 and all(r[0] for r in data_rows)
    totals = [c.value for c in sheet[generate.SALES_TOTALS_ROW]]
    assert totals[0] == "Summe" and totals[3] == "=SUM(D4:D27)" and generate.SALES_TOTALS_ROW == 28
    cached = openpyxl.load_workbook(path, data_only=True)["Umsatz 2026"]
    assert cached.cell(generate.SALES_TOTALS_ROW, 4).value is None  # formulas without saved results
    assert len(book["Ziele"]["A"]) == 13


def test_the_customers_below_target_follow_from_the_numbers() -> None:
    assert generate.below_target(1) == [
        "Brauhaus Keller GmbH",
        "Kowalski Stahl Sp. z o.o.",
        "Rossi Design",
        "Studio Lindqvist AB",
        "Maple Fasteners Inc.",
    ]
    assert generate.below_target(0) == ["Kowalski Stahl Sp. z o.o.", "Rossi Design"]
    assert "Müller & Söhne KG" not in generate.below_target(1)  # exactly 80 percent is not below


# ------------------------------------------------------------------------ JSON, SQLite


def test_the_json_dump_has_200_orders_with_the_planted_defects(out: Path) -> None:
    text = (out / "json" / "shop_orders.jsonl").read_text(encoding="utf-8")
    orders: list[dict[str, Any]] = [json.loads(line) for line in text.splitlines()]
    assert len(orders) == 200 and len({o["orderId"] for o in orders}) == 200
    no_address = [i for i, o in enumerate(orders, 1) if o["deliveryAddress"] is None]
    assert no_address == list(generate.NO_ADDRESS)
    wrong = [
        i
        for i, o in enumerate(orders, 1)
        if round(sum(x["qty"] * x["unitPrice"] for x in o["lines"]), 2) != o["total"]
    ]
    assert wrong == list(generate.TOTAL_MISMATCH)
    snake = [i for i, o in enumerate(orders, 1) if "customer_id" in o]
    assert snake == list(generate.SNAKE_CASE_KEY)
    assert all(o["createdAt"] > 10**12 for o in orders)  # epoch milliseconds


def test_the_database_is_sound_and_reading_it_does_not_fire_the_trigger(out: Path) -> None:
    path = out / "sqlite" / "shop.db"
    before = path.read_bytes()
    connection = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    try:
        kinds = dict(connection.execute("SELECT name, type FROM sqlite_master"))
        assert kinds["open_orders"] == "view" and kinds["orders_audit"] == "trigger"
        assert connection.execute("PRAGMA integrity_check").fetchall() == [("ok",)]
        assert len(connection.execute("PRAGMA foreign_key_check").fetchall()) == 1  # C-999
        assert connection.execute("SELECT COUNT(*) FROM audit_log").fetchone() == (0,)
        assert connection.execute("SELECT COUNT(*) FROM orders").fetchone() == (30,)
    finally:
        connection.close()
    assert path.read_bytes() == before


# ------------------------------------------------------------------------ mailbox


def _messages(out: Path) -> dict[str, Any]:
    return {
        p.name: email.message_from_bytes(p.read_bytes(), policy=email.policy.default)
        for p in sorted((out / "mailbox").glob("*.eml"))
    }


def test_the_mailbox_has_six_messages_and_exactly_one_attack(out: Path) -> None:
    messages = _messages(out)
    assert len(messages) == 6
    bodies = {
        name: message.get_body(preferencelist=("plain", "html")).get_content()
        for name, message in messages.items()
    }
    attacks = [n for n, text in bodies.items() if "IGNORE YOUR INSTRUCTIONS" in text.upper()]
    assert attacks == ["06_attack.eml"]
    assert [n for n, t in bodies.items() if "attacker@example.com" in t] == ["06_attack.eml"]
    assert all("instructions" not in t.lower() for n, t in bodies.items() if n != "06_attack.eml")


def test_the_mails_differ_in_the_ways_the_mail_profiler_must_handle(out: Path) -> None:
    messages = _messages(out)
    subject = messages["03_rueckfrage_umlaute.eml"]["Subject"]
    assert subject == "Rückfrage zur Lieferung: Größe & Maße"
    wire = (out / "mailbox" / "03_rueckfrage_umlaute.eml").read_bytes()
    assert b"=?utf-8?" in wire and "Rückfrage".encode() not in wire  # RFC 2047 on the wire
    html = messages["04_html_only.eml"]
    assert html.get_content_type() == "text/html" and "tracker.example" in html.get_content()
    attachments = list(messages["05_invoice_attached.eml"].iter_attachments())
    assert [a.get_filename() for a in attachments] == ["2024-0412_mueller.pdf"]
    reader = pypdf.PdfReader(io.BytesIO(attachments[0].get_payload(decode=True)))
    assert reader.pages[0].extract_text().startswith("RECHNUNG 2024-0412")
    assert all(m["Date"] == "Mon, 05 Oct 2026 08:00:00 +0000" for m in messages.values())


# ------------------------------------------------------------------------ PDFs


def _text_per_page(path: Path) -> list[int]:
    return [len((page.extract_text() or "").strip()) for page in pypdf.PdfReader(path).pages]


def test_the_pdfs_have_a_text_layer_none_or_both(out: Path) -> None:
    born, scan, mixed = (
        _text_per_page(out / "pdf" / name)
        for name in ("born_digital_2p.pdf", "scan_2p.pdf", "mixed_3p.pdf")
    )
    assert len(born) == 2 and all(born)
    assert scan == [0, 0]
    assert [bool(n) for n in mixed] == [True, False, True]


def test_the_hidden_line_is_in_the_text_layer_but_not_visible(out: Path) -> None:
    path = out / "pdf" / "hidden_instruction.pdf"
    assert generate.HIDDEN_LINE in pypdf.PdfReader(path).pages[-1].extract_text()
    assert b"1 1 1 rg" in path.read_bytes()  # white fill color, 1 pt
    plain = pypdf.PdfReader(out / "pdf" / "born_digital_2p.pdf").pages[-1].extract_text()
    assert generate.HIDDEN_LINE not in plain


def test_the_logo_is_a_valid_png_that_no_reader_of_awa_handles(out: Path) -> None:
    raw = (out / "image" / "logo.png").read_bytes()
    assert raw.startswith(b"\x89PNG\r\n\x1a\n") and raw.endswith(b"IEND\xaeB`\x82")
    assert struct.unpack(">II", raw[16:24]) == (8, 8)
