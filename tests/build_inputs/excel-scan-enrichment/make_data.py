"""Generate the test input for a from-scratch build: an invoice register and scanned invoices.

Test input only, not a solution: `awa build` gets `process.md` and the generated `data/`
folder and must design the system itself. Everything is fictitious and deterministic.

    python tests/build_inputs/excel-scan-enrichment/make_data.py [--out DIR] [--rows N]

Writes into DIR (default: `data/` beside this file, which git ignores):

- `register.xlsx`: sheet `Incoming invoices` (N rows, default 4000; formatted header, an
  autofilter, a frozen pane) and a sheet `Cost centers`;
- `scans/*.pdf`: 30 image-only PDFs (one JPEG per page, no text layer): clean, rotated,
  faint, noisy, one without a VAT id, one whose invoice number is in no row, three with German
  number formats, one with spaced key;
- `ORIGIN.md`;
- `expected.json` (beside DIR, not in it, so the build never sees it): what a perfect run
  writes per invoice; `check_output.py` scores an output workbook against it.

Needs Pillow and openpyxl.
"""

from __future__ import annotations

import argparse
import json
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFilter, ImageFont

HERE = Path(__file__).resolve().parent
SUPPLIERS = [
    ("Nordlicht Bürobedarf GmbH", "DE"),
    ("Kessler Maschinenteile KG", "DE"),
    ("Alpenblick Logistik AG", "AT"),
    ("Rheinufer Druckerei GmbH", "DE"),
    ("Brightwater Components Ltd", "IE"),
    ("Lindgren Verktyg AB", "SE"),
    ("Sonnenfeld Elektro GmbH", "DE"),
    ("Van der Berg Verpakking BV", "NL"),
]
TERMS = ["30 days net", "14 days 2% discount, 30 days net", "10 days net", "60 days net"]
COST_CENTERS = ["CC-100 Administration", "CC-200 Production", "CC-300 Logistics", "CC-400 Sales"]
COLUMNS = ["Invoice no", "Supplier", "Received", "Gross amount (EUR)", "Cost center", "Status"]


def font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for name in ("DejaVuSans.ttf", "arial.ttf", "Arial.ttf", "LiberationSans-Regular.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def vat_id(country: str, rng: random.Random) -> str:
    return country + "".join(str(rng.randint(0, 9)) for _ in range(9))


def money(value: float, german: bool) -> str:
    text = f"{value:,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".") if german else text


def page(fields: dict[str, object], kind: str, rng: random.Random) -> Image.Image:
    """An A4-like invoice picture (1240 x 1754, 150 dpi), made messy by `kind`."""

    image = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(image)
    big, normal = font(44), font(28)
    german = kind == "german"
    supplier = str(fields["supplier"])
    draw.text((90, 90), supplier, font=big, fill=0)
    draw.text((90, 150), "Industriestrasse 12, 12345 Musterstadt", font=normal, fill=60)
    if fields.get("vat_id"):
        label = "USt-IdNr." if german else "VAT ID"
        draw.text((90, 190), f"{label}: {fields['vat_id']}", font=normal, fill=0)
    draw.text((90, 300), "Rechnung" if german else "INVOICE", font=big, fill=0)
    number = str(fields["invoice_no"])
    if kind == "spaced":
        number = number.replace("-", " ").lower()
    issued: date = fields["invoice_date"]  # type: ignore[assignment]
    shown = issued.strftime("%d.%m.%Y") if german else issued.strftime("%d %b %Y")
    rows = [
        ("Rechnungsnr." if german else "Invoice No.", number),
        ("Rechnungsdatum" if german else "Invoice date", shown),
        ("Kunde" if german else "Customer", "Example Manufacturing GmbH"),
    ]
    for index, (label, value) in enumerate(rows):
        draw.text((90, 400 + index * 50), f"{label}:", font=normal, fill=0)
        draw.text((480, 400 + index * 50), value, font=normal, fill=0)
    draw.line((90, 600, 1150, 600), fill=0, width=2)
    draw.text((90, 620), "Description", font=normal, fill=0)
    draw.text((900, 620), "Amount EUR", font=normal, fill=0)
    draw.text((90, 680), "Goods and services as per order", font=normal, fill=0)
    draw.text((900, 680), money(float(fields["net"]), german), font=normal, fill=0)  # type: ignore[arg-type]
    draw.line((90, 760, 1150, 760), fill=0, width=1)
    totals = [
        ("Nettobetrag" if german else "Net amount", float(fields["net"])),  # type: ignore[arg-type]
        (f"MwSt {fields['rate']}%" if german else f"VAT {fields['rate']}%", float(fields["vat"])),  # type: ignore[arg-type]
        ("Gesamtbetrag" if german else "Total", float(fields["gross"])),  # type: ignore[arg-type]
    ]
    for index, (label, value) in enumerate(totals):
        draw.text((600, 790 + index * 50), label, font=normal, fill=0)
        draw.text((900, 790 + index * 50), money(value, german), font=normal, fill=0)
    terms_label = "Zahlungsbedingungen" if german else "Payment terms"
    draw.text((90, 1050), f"{terms_label}: {fields['terms']}", font=normal, fill=0)
    draw.text(
        (90, 1650), "Bank: Example Bank, IBAN DE00 0000 0000 0000 0000 00", font=normal, fill=90
    )
    if kind == "faint":
        image = image.point(lambda v: 150 + v * 105 // 255)
    if kind in ("noisy", "faint"):
        noise = Image.effect_noise(image.size, 40)
        image = Image.blend(image, noise, 0.18)
    image = image.rotate(rng.uniform(-1.5, 1.5), fillcolor=255, resample=Image.BICUBIC)
    if kind == "rotated90":
        image = image.rotate(90, expand=True)
    elif kind == "rotated180":
        image = image.rotate(180)
    return image.filter(ImageFilter.GaussianBlur(0.6))


SCAN_KINDS = (
    ["clean"] * 14
    + ["rotated90"] * 2
    + ["rotated180"]
    + ["faint"] * 3
    + ["noisy"] * 4
    + ["german"] * 3
    + ["spaced"]
    + ["no_vat_id"]
)
MIN_ROWS = len(SCAN_KINDS) + 1  # one register row per scan, plus the unmatched scan


def build(out: Path, rows: int) -> dict[str, object]:
    rng = random.Random(20261008)
    out.mkdir(parents=True, exist_ok=True)
    (out / "scans").mkdir(exist_ok=True)
    for old in (out / "scans").glob("*.pdf"):
        old.unlink()
    vat_ids = {name: vat_id(country, rng) for name, country in SUPPLIERS}
    book = Workbook()
    sheet = book.active
    sheet.title = "Incoming invoices"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    start = date(2026, 1, 2)
    register: list[dict[str, object]] = []
    for number in range(1, rows + 1):
        supplier = rng.choice(SUPPLIERS)[0]
        received = start + timedelta(days=number * 200 // rows)
        net = round(rng.uniform(80, 25000), 2)
        rate = 19 if rng.random() < 0.85 else 7
        vat = round(net * rate / 100, 2)
        row = {
            "invoice_no": f"INV-2026-{number:05d}",
            "supplier": supplier,
            "received": received,
            "net": net,
            "rate": rate,
            "vat": vat,
            "gross": round(net + vat, 2),
            "terms": rng.choice(TERMS),
            "invoice_date": received - timedelta(days=rng.randint(1, 6)),
            "vat_id": vat_ids[supplier],
        }
        register.append(row)
        sheet.append(
            [row["invoice_no"], supplier, received, row["gross"], rng.choice(COST_CENTERS), "open"]
        )
        sheet.cell(sheet.max_row, 3).number_format = "dd.mm.yyyy"
        sheet.cell(sheet.max_row, 4).number_format = "#,##0.00"
    sheet.auto_filter.ref = f"A1:F{rows + 1}"
    sheet.freeze_panes = "A2"
    for letter, width in zip("ABCDEF", (18, 30, 12, 20, 24, 10), strict=True):
        sheet.column_dimensions[letter].width = width
    centers = book.create_sheet("Cost centers")
    centers.append(["Cost center", "Owner"])
    for index, center in enumerate(COST_CENTERS):
        centers.append([center, f"Owner {index + 1}"])
    book.save(out / "register.xlsx")

    chosen = rng.sample(register, MIN_ROWS)
    expected: dict[str, object] = {}
    for index, (kind, row) in enumerate(zip(SCAN_KINDS, chosen, strict=False)):
        fields = dict(row)
        if kind == "no_vat_id":
            fields["vat_id"] = None
        name = f"scan_{index + 1:03d}.pdf"
        page(fields, kind, rng).convert("RGB").save(
            out / "scans" / name, "PDF", resolution=150, quality=80
        )
        expected[name] = {
            "kind": kind,
            "invoice_no": row["invoice_no"],
            "invoice_date": row["invoice_date"].isoformat(),  # type: ignore[attr-defined]
            "net": row["net"],
            "vat": row["vat"],
            "vat_id": fields["vat_id"],
            "terms": row["terms"],
        }
    orphan = dict(chosen[-1])
    orphan["invoice_no"] = "INV-2026-99999"
    page(orphan, "clean", rng).convert("RGB").save(
        out / "scans" / "scan_030.pdf", "PDF", resolution=150, quality=80
    )
    expected["scan_030.pdf"] = {"kind": "unmatched", "invoice_no": "INV-2026-99999"}
    (out / "ORIGIN.md").write_text(
        "Test data: public or generated, not company data.\n\n"
        "Generated by tests/build_inputs/excel-scan-enrichment/make_data.py (fixed seed); all "
        "suppliers, numbers and amounts are fictitious. The scans are image-only PDFs without "
        "a text layer.\n",
        encoding="utf-8",
    )
    return expected


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=HERE / "data")
    parser.add_argument("--rows", type=int, default=4000)
    args = parser.parse_args()
    if args.rows < MIN_ROWS:
        parser.error(f"--rows must be at least {MIN_ROWS}, one per scan")
    expected = build(args.out, args.rows)
    target = args.out.parent / "expected.json"
    target.write_text(json.dumps(expected, indent=1), encoding="utf-8")
    print(f"wrote {args.out} ({args.rows} rows, {len(expected)} scans) and {target}")


if __name__ == "__main__":
    main()
