"""Score an enriched register against `expected.json` (written by `make_data.py`).

    python tests/build_inputs/excel-scan-enrichment/check_output.py OUTPUT.xlsx

Finds the new columns by their names (date, net, VAT amount, VAT id, terms; any wording),
then compares, per scanned invoice, the five values in its row: dates as dates, amounts to
the cent, the VAT id ignoring spaces, terms ignoring case and punctuation. Prints the field
and invoice pass rates and whether the unmatched invoice was kept.
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
FIELDS = {
    "invoice_date": re.compile(r"date", re.I),
    "net": re.compile(r"net", re.I),
    "vat": re.compile(r"(vat|tax|mwst).*(amount|eur)|^(vat|tax|mwst)( amount)?$", re.I),
    "vat_id": re.compile(r"(vat|tax|ust).*id", re.I),
    "terms": re.compile(r"term", re.I),
}


def norm(text: object) -> str:
    return re.sub(r"[\W_]+", "", str(text or "")).upper()


def number(value: object) -> float | None:
    if isinstance(value, int | float):
        return float(value)
    text = re.sub(r"[^\d,.\-]", "", str(value or ""))
    if not text:
        return None
    if "," in text and text.rfind(",") > text.rfind("."):
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d %b %Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def same(field: str, want: object, got: object) -> bool:
    if want is None:
        return got in (None, "")
    if field == "invoice_date":
        return as_date(got) == date.fromisoformat(str(want))
    if field in ("net", "vat"):
        found = number(got)
        return found is not None and abs(found - float(want)) < 0.006  # type: ignore[arg-type]
    return norm(got) == norm(want)


def main(path: str) -> int:
    expected = json.loads((HERE / "expected.json").read_text(encoding="utf-8"))
    book = load_workbook(path, data_only=True)
    sheet = book["Incoming invoices"] if "Incoming invoices" in book.sheetnames else book.active
    header = [str(c.value or "") for c in sheet[1]]
    added = header[6:]
    columns: dict[str, int] = {}
    for field, pattern in FIELDS.items():
        for index, name in enumerate(added, start=6):
            if pattern.search(name) and index not in columns.values():
                if field == "net" and re.search(r"vat|tax", name, re.I):
                    continue
                columns[field] = index
                break
    print(f"new columns: {added}; mapped: { {k: header[v] for k, v in columns.items()} }")
    rows = {norm(r[0].value): r for r in sheet.iter_rows(min_row=2)}
    fields_ok = fields_all = invoices_ok = 0
    failures: list[str] = []
    for scan, want in expected.items():
        if want["kind"] == "unmatched":
            continue
        row = rows.get(norm(want["invoice_no"]))
        good = True
        for field in FIELDS:
            fields_all += 1
            got = row[columns[field]].value if row is not None and field in columns else None
            if same(field, want[field], got):
                fields_ok += 1
            else:
                good = False
                failures.append(
                    f"{scan} ({want['kind']}) {field}: want {want[field]!r} got {got!r}"
                )
        invoices_ok += good
    total = sum(1 for w in expected.values() if w["kind"] != "unmatched")
    orphan = any(
        "99999" in str(cell.value or "")
        for name in book.sheetnames
        if name != sheet.title
        for row in book[name].iter_rows()
        for cell in row
    )
    print(f"fields: {fields_ok}/{fields_all} = {fields_ok / fields_all:.1%}")
    print(f"invoices fully right: {invoices_ok}/{total}")
    print(f"unmatched invoice kept elsewhere: {orphan}")
    for line in failures:
        print("  " + line)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
