"""Relations between files and the flag catalog (`awa.data._relations`, `_values`, `_flags`)."""

from __future__ import annotations

import ast
import json
import re
import sqlite3
import time
from collections.abc import Mapping, Sequence
from pathlib import Path

import jsonschema
import pytest
from openpyxl import Workbook

from _helpers.data_fixtures import REPO
from _helpers.data_workbooks import rows_to, save
from awa.data import (
    FLAGS,
    ColumnValues,
    DataProfile,
    Relation,
    fix_hint,
    make_flag,
    profile,
    relations,
)
from awa.data._model import FileProfile
from awa.data._relations import _duplicate_files, analyze_links, with_relations
from awa.data._values import local_values

DATA = REPO / "examples/usecases"
O2C = DATA / "order-to-cash-exceptions/data"


def write(folder: Path, name: str, text: str, encoding: str = "utf-8") -> Path:
    path = folder / name
    path.write_bytes(text.encode(encoding))
    return path


def csv(header: str, *rows: str) -> str:
    return "\n".join([header, *rows]) + "\n"


def link(found: Sequence[Relation], source: str, target: str) -> Relation | None:
    return next((r for r in found if (r.from_column, r.to_column) == (source, target)), None)


def codes(flags: Sequence[object]) -> list[str]:
    return [f.code for f in flags]  # type: ignore[attr-defined]


# --- joins between files -----------------------------------------------------------------------


def test_orders_point_at_their_customers_with_full_overlap(tmp_path: Path) -> None:
    write(
        tmp_path,
        "customers.csv",
        csv("customer_id;name", *[f"C-{n};Firma {n}" for n in range(10, 70, 10)]),
    )
    write(
        tmp_path,
        "orders.csv",
        csv(
            "order_id;customer_id;amount", *[f"O-{n};C-{10 * (n % 6 + 1)};{n}" for n in range(1, 9)]
        ),
    )

    p = profile(tmp_path)
    found = link(p.relations, "f2.t1.customer_id", "f1.t1.customer_id")

    assert found is not None and found.kind == "value_overlap" and found.overlap == 1.0
    assert found.confidence >= 0.8
    assert not p.flags  # a clean join raises nothing


def test_a_reference_column_that_quotes_the_invoice_numbers_is_a_weak_link_with_a_warning() -> None:
    p = profile(O2C)
    invoices = p.find_column("invoices.xlsx", "invoice_no")
    reference = p.find_column("payments.csv", "reference")
    assert invoices is not None and reference is not None
    source = f"{p.find_table('payments.csv').id}.reference"  # type: ignore[union-attr]
    target = f"{p.find_table('invoices.xlsx').id}.invoice_no"  # type: ignore[union-attr]

    found = link(p.relations, source, target)

    assert found is not None and found.overlap == pytest.approx(0.3333, abs=1e-4)
    assert found.confidence < 0.5  # only a link above 0.8 may be used without asking
    flag = next(f for f in p.flags if f.code == "DF-W-REFERENCE-NEEDS-PARSING")
    assert flag.where == source and "3 of 6 distinct values contain a key" in flag.message


def test_the_warehouse_database_and_the_erp_export_share_the_order_id(tmp_path: Path) -> None:
    folder = tmp_path / "in"
    folder.mkdir()
    for name in ("warehouse.db",):
        (folder / name).write_bytes((DATA / "data-quality-reconciliation/data" / name).read_bytes())
    export = DATA / "data-quality-reconciliation/data/samples/erp_export_2026-10-01.xlsx"
    (folder / export.name).write_bytes(export.read_bytes())

    p = profile(folder)
    pairs = [
        r
        for r in p.relations
        if r.from_column.endswith(".order_id") and r.to_column.endswith(".order_id")
    ]

    (pair,) = pairs
    assert (
        pair.kind == "value_overlap"
        and pair.confidence >= 0.8
        and pair.overlap == pytest.approx(0.9672, abs=1e-4)
    )
    orphan = next(f for f in p.flags if f.code == "DF-W-ORPHAN-VALUES")
    assert orphan.where == pair.from_column and "2 of 61 distinct values" in orphan.message
    # the values themselves are samples: in the relation, never in a message that a privacy
    # mode would still print
    assert pair.unmatched == ["O-1061", "O-1062"]
    assert "O-1061" not in orphan.message and "O-1062" not in orphan.message
    stripped = p.without_samples()
    assert "O-1061" not in stripped.to_json() and "O-1062" not in stripped.to_json()
    assert [r.confidence for r in stripped.relations] == [r.confidence for r in p.relations]


def test_keys_that_differ_only_in_spelling_are_joined_and_flagged(tmp_path: Path) -> None:
    write(
        tmp_path,
        "invoices.csv",
        csv("invoice_no;amount", *[f"INV-{n};{n}" for n in range(3001, 3009)]),
    )
    write(
        tmp_path,
        "payments.csv",
        csv(
            "invoice_no;paid",
            "INV-3001;1",
            "INV-3002;2",
            "inv 3003;3",
            "Inv-3004;4",
            "INV-3005;5",
            "INV-3006;6",
            "INV-9999;7",
        ),
    )

    p = profile(tmp_path)
    found = link(p.relations, "f2.t1.invoice_no", "f1.t1.invoice_no")

    assert (
        found is not None
        and found.overlap == pytest.approx(6 / 7, abs=1e-4)
        and found.confidence >= 0.8
    )
    assert {"DF-W-KEY-SPELLING-DIFFERS", "DF-W-ORPHAN-VALUES"} <= set(codes(p.flags))
    spelling = next(f for f in p.flags if f.code == "DF-W-KEY-SPELLING-DIFFERS")
    assert "2 of 7 distinct values" in spelling.message


def test_a_key_quoted_inside_text_is_a_weak_link_but_a_bare_number_is_not_enough(
    tmp_path: Path,
) -> None:
    invoices = csv("invoice_no;amount", *[f"INV-{n};{n}" for n in range(3001, 3007)])
    write(tmp_path, "invoices.csv", invoices)
    write(
        tmp_path, "memos.csv", csv("memo;x", *[f"INV-{n} Mueller;{n}" for n in range(3001, 3007)])
    )

    p = profile(tmp_path)
    found = link(p.relations, "f2.t1.memo", "f1.t1.invoice_no")

    # the whole key stands in the text: it counts 0.6 and a link is made, never a sure one
    assert found is not None and found.overlap == 0.0 and found.confidence < 0.5
    flag = next(f for f in p.flags if f.code == "DF-W-REFERENCE-NEEDS-PARSING")
    assert "6 of 6 distinct values contain a key" in flag.message

    write(tmp_path, "memos.csv", csv("memo;x", *[f"RE {n};{n}" for n in range(3001, 3007)]))
    # only the number of the key under another prefix counts 0.4: below the half that is needed
    assert link(profile(tmp_path).relations, "f2.t1.memo", "f1.t1.invoice_no") is None


def test_a_short_row_does_not_hide_the_values_of_the_other_rows(tmp_path: Path) -> None:
    write(
        tmp_path,
        "customers.csv",
        csv("customer_id;name", *[f"C-{n};Firma {n}" for n in range(10, 70, 10)]),
    )
    rows = [f"O-{n};{n};C-{10 * (n % 6 + 1)}" for n in range(1, 9)]
    write(tmp_path, "orders.csv", csv("order_id;amount;customer_id", *rows, "O-9;9"))

    found = link(profile(tmp_path).relations, "f2.t1.customer_id", "f1.t1.customer_id")

    assert found is not None and found.kind == "value_overlap" and found.overlap == 1.0


def test_equal_names_with_disjoint_values_are_not_a_relation(tmp_path: Path) -> None:
    write(tmp_path, "a.csv", csv("customer;x", *[f"A{n};{n}" for n in range(8)]))
    write(tmp_path, "b.csv", csv("customer;y", *[f"B{n};{n}" for n in range(8)]))

    assert profile(tmp_path).relations == []


def test_line_numbers_and_quantities_are_no_keys_but_integer_ids_are(tmp_path: Path) -> None:
    write(
        tmp_path, "r1.csv", csv("line;qty;customer_id", *[f"{n};{n + 1};{n}" for n in range(1, 8)])
    )
    write(
        tmp_path, "r2.csv", csv("line;qty;customer_id", *[f"{n};{n + 1};{n}" for n in range(1, 8)])
    )

    pairs = {
        (r.from_column.rpartition(".")[2], r.to_column.rpartition(".")[2])
        for r in profile(tmp_path).relations
    }

    assert pairs == {("customer_id", "customer_id")}


def test_free_text_and_dates_are_never_compared(tmp_path: Path) -> None:
    text = "a long free text that goes on and on and on and on and on and on and on and on"
    write(tmp_path, "a.csv", csv("note;day", *[f"{text} {n};2024-01-{n:02d}" for n in range(1, 8)]))
    write(tmp_path, "b.csv", csv("note;day", *[f"{text} {n};2024-01-{n:02d}" for n in range(1, 8)]))

    assert profile(tmp_path).relations == []


def test_the_same_file_twice_is_reported(tmp_path: Path) -> None:
    write(tmp_path, "a.csv", csv("x;y", "1;a", "2;b"))
    write(tmp_path, "copy of a.csv", csv("x;y", "1;a", "2;b"))

    p = profile(tmp_path)

    flag = next(f for f in p.flags if f.code == "DF-W-DUPLICATE-FILE")
    assert flag.where == "f2" and "exactly the content of a.csv" in flag.message


def test_unreadable_files_are_no_duplicates_of_each_other() -> None:
    files = [
        FileProfile(id=f"f{n}", path=f"{n}.pdf", sha256="", size=10, kind="pdf", status="error")
        for n in (1, 2)
    ]

    assert _duplicate_files(DataProfile(awa_version="1", created="x", files=files)) == []


# --- what is declared, what is only a name ------------------------------------------------------------


def test_a_declared_foreign_key_is_certain_and_not_repeated_as_a_guess(tmp_path: Path) -> None:
    path = tmp_path / "shop.db"
    connection = sqlite3.connect(path)
    connection.executescript(
        """
        CREATE TABLE customers (customer_id INTEGER PRIMARY KEY, name TEXT);
        CREATE TABLE orders (order_id INTEGER PRIMARY KEY, customer_id INTEGER REFERENCES customers(customer_id));
        INSERT INTO customers VALUES (1,'a'),(2,'b'),(3,'c'),(4,'d'),(5,'e');
        INSERT INTO orders VALUES (1,1),(2,2),(3,2),(4,3),(5,4),(6,5);
        """
    )
    connection.commit()
    connection.close()

    p = profile(path)

    (fk,) = p.relations  # inside one file too: every relation is in DataProfile.relations
    assert (fk.kind, fk.confidence, fk.overlap) == ("foreign_key", 1.0, None)
    assert fk.from_column.endswith(".customer_id") and fk.to_column.endswith(".customer_id")
    assert p.files[0].relations == []


def test_without_values_a_name_alone_never_reaches_the_threshold(tmp_path: Path) -> None:
    write(tmp_path, "customers.csv", csv("id;name", *[f"{n};x{n}" for n in range(1, 8)]))
    write(
        tmp_path,
        "orders.csv",
        csv("order_id;customers_id;amount", *[f"{n};{n};{n}" for n in range(1, 8)]),
    )
    p = profile(tmp_path)

    blind = relations(p, values=lambda ids: {})

    assert blind and all(r.kind == "name_match" and r.confidence < 0.8 for r in blind)
    assert {(r.from_column.rpartition(".")[2], r.to_column.rpartition(".")[2]) for r in blind} == {
        ("customers_id", "id")
    }


def test_the_value_source_is_a_port_and_is_asked_only_for_candidate_columns(tmp_path: Path) -> None:
    write(tmp_path, "a.csv", csv("customer_id;memo", *[f"C-{n};m{n}" for n in range(8)]))
    write(tmp_path, "b.csv", csv("customer_id;other", *[f"C-{n};o{n}" for n in range(8)]))
    p = profile(tmp_path)
    asked: list[list[str]] = []

    def source(ids: Sequence[str]) -> Mapping[str, ColumnValues]:
        asked.append(list(ids))
        return {i: ColumnValues(frozenset({f"C-{n}" for n in range(8)}), 8, True) for i in ids}

    found = relations(p, values=source)

    assert asked == [["f1.t1.customer_id", "f2.t1.customer_id"]]  # not the memo columns
    assert [(r.overlap, r.kind) for r in found] == [(1.0, "value_overlap")]
    asked.clear()
    relations(profile(write(tmp_path, "lone.csv", "x\n1\n")), values=source)
    assert asked == []  # one table: no pairs, no reading


def test_the_same_profile_gives_the_same_relations_and_the_json_validates() -> None:
    p = profile(O2C, created="2026-10-05T00:00:00Z")

    again = profile(O2C, created="2026-10-05T00:00:00Z")

    assert p.to_json() == again.to_json() and p.relations
    jsonschema.validate(
        json.loads(p.to_json()), DataProfile.json_schema(), jsonschema.Draft202012Validator
    )
    order = [(-r.confidence, r.from_column, r.to_column) for r in p.relations]
    assert order == sorted(order)  # strongest first, ties by id


def test_many_files_with_the_same_column_names_stay_bounded_and_fast(tmp_path: Path) -> None:
    for number in range(70):
        write(
            tmp_path,
            f"f{number:02d}.csv",
            csv("id;name;status", *[f"ID-{n};name {n};open" for n in range(20)]),
        )
    started = time.perf_counter()

    p = profile(tmp_path)

    assert len(p.relations) <= 400 and time.perf_counter() - started < 60


def test_links_between_sheets_and_between_files_are_all_in_the_profile(tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Kunden"
    rows_to(sheet, [["kundennr", "name"], *[[n, f"k{n}"] for n in range(1, 8)]])
    other = book.create_sheet("Bestellungen")
    rows_to(other, [["bestellnr", "kundennr"], *[[100 + n, n] for n in range(1, 8)]])
    save(book, tmp_path, "shop.xlsx")
    write(tmp_path, "tickets.csv", csv("ticket;kundennr", *[f"T-{n};{n}" for n in range(1, 8)]))

    p = profile(tmp_path)

    pairs = {(r.from_column, r.to_column) for r in p.relations}
    assert ("f1.t2.kundennr", "f1.t1.kundennr") in pairs  # sheet to sheet
    assert any(a.startswith("f2.") and b.startswith("f1.") for a, b in pairs)  # file to file
    assert p.summary.warnings == sum(
        f.severity == "warning" for f in [*p.flags, *(x for fl in p.files for x in fl.flags)]
    )


# --- reading the values again ------------------------------------------------------------------------------


def test_local_values_follow_the_decisions_of_the_profile(tmp_path: Path) -> None:
    write(
        tmp_path,
        "bank.csv",
        csv(
            "Umsatzanzeige;Konto",
            "Buchungstag;Betrag;Text",
            "02.09.2024;1,5;Müller",
            "03.09.2024;2,5;Köhler",
            "04.09.2024;;Müller",
            "Summe;4,0;",
        ),
        encoding="cp1252",
    )
    p = profile(tmp_path / "bank.csv")
    (file,) = p.files
    table = file.tables[0]
    read = local_values(p)

    values = read([f"{table.id}.text", f"{table.id}.betrag", f"{table.id}.missing"])

    assert set(values) == {f"{table.id}.text", f"{table.id}.betrag"}
    assert values[f"{table.id}.text"] == ColumnValues(frozenset({"Müller", "Köhler"}), 3, True)
    # the totals row is not data
    assert values[f"{table.id}.betrag"].values == frozenset({"1,5", "2,5"})


def test_local_values_of_a_workbook_table_that_does_not_start_in_column_a(tmp_path: Path) -> None:
    book = Workbook()
    rows_to(
        book.active,
        [["Titel"], [], ["Name", "Qty"], ["x", 1], ["y", 2], ["Summe", "=SUM(B4:B5)"]],
        start=1,
        column=3,
    )
    save(book, tmp_path, "w.xlsx")
    p = profile(tmp_path / "w.xlsx")
    table = p.files[0].tables[0]

    values = local_values(p)([f"{table.id}.name", f"{table.id}.qty"])

    assert table.range == "C3:D6"
    assert values[f"{table.id}.name"].values == frozenset({"x", "y"})
    assert values[f"{table.id}.qty"].values == frozenset({"1", "2"})


def test_local_values_never_raise_for_files_that_are_gone_or_hostile(tmp_path: Path) -> None:
    write(tmp_path, "a.csv", csv("x;y", "1;a", "2;b"))
    p = profile(tmp_path / "a.csv")
    table = p.files[0].tables[0]
    ids = [f"{table.id}.x", "f9.t9.nothing", "garbage"]
    (tmp_path / "a.csv").unlink()

    assert local_values(p)(ids) == {}
    assert local_values(DataProfile(awa_version="1", created="x"))(["f1.t1.x"]) == {}


def test_a_file_that_changed_since_profiling_is_not_read(tmp_path: Path) -> None:
    path = write(tmp_path, "a.csv", csv("x;y", "1;a", "2;b"))
    p = profile(path)
    table = p.files[0].tables[0]
    path.write_bytes(path.read_bytes() + b"3;c\n")

    assert local_values(p)([f"{table.id}.x"]) == {}


def test_local_values_are_bounded(tmp_path: Path) -> None:
    write(tmp_path, "big.csv", csv("id;grp", *[f"{n};g{n % 3}" for n in range(500)]))
    p = profile(tmp_path / "big.csv")
    table = p.files[0].tables[0]

    values = local_values(p, max_rows=100, max_distinct=10)([f"{table.id}.id", f"{table.id}.grp"])

    assert not values[f"{table.id}.id"].complete and len(values[f"{table.id}.id"].values) == 10
    assert not values[f"{table.id}.grp"].complete and values[f"{table.id}.grp"].cells == 100


def test_local_values_never_modify_the_database_and_quote_odd_names(tmp_path: Path) -> None:
    path = tmp_path / "odd.db"
    connection = sqlite3.connect(path)
    connection.executescript(
        'CREATE TABLE "we ""ird" ("a b" TEXT, "c""d" TEXT); INSERT INTO "we ""ird" VALUES (\'x\',\'y\');'
    )
    connection.commit()
    connection.close()
    before = path.read_bytes()
    p = profile(path)
    table = p.files[0].tables[0]

    values = local_values(p)([f"{table.id}.{c.key}" for c in table.columns])

    assert sorted(v.values for v in values.values()) == [frozenset({"x"}), frozenset({"y"})]
    assert path.read_bytes() == before


# --- the flag catalog ------------------------------------------------------------------------------------------


def _codes_in_source() -> dict[str, set[str]]:
    pattern = re.compile(r"^DF-[IWE]-[A-Z0-9]+(?:-[A-Z0-9]+)*$")
    found: dict[str, set[str]] = {}
    for path in sorted((REPO / "src/awa/data").glob("*.py")):
        if path.name == "_flags.py":
            continue
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if (
                isinstance(node, ast.Constant)
                and isinstance(node.value, str)
                and pattern.match(node.value)
            ):
                found.setdefault(node.value, set()).add(path.name)
    return found


def test_every_flag_code_in_the_code_is_in_the_catalog_and_every_entry_is_used() -> None:
    used = _codes_in_source()

    assert set(used) - set(FLAGS) == set(), "add a fix hint to _flags.FLAGS"
    assert set(FLAGS) - set(used) == set(), "remove the entry or use the code"


def test_the_catalog_is_well_formed() -> None:
    for code, hint in FLAGS.items():
        assert re.fullmatch(r"DF-[IWE]-[A-Z0-9]+(?:-[A-Z0-9]+)*", code), code
        assert hint.endswith((".", ")")) and 25 <= len(hint) <= 200, code
    assert fix_hint("DF-W-DECIMAL-COMMA") == FLAGS["DF-W-DECIMAL-COMMA"]
    assert "No fix hint" in fix_hint("DF-W-COMPANY-OWN")
    with pytest.raises(TypeError):
        FLAGS["DF-W-NEW"] = "x"  # type: ignore[index]


def test_a_flag_of_every_family_is_made_with_the_severity_of_its_family() -> None:
    severities = {code[3]: make_flag(code, "x").severity for code in FLAGS}

    assert severities == {"I": "info", "W": "warning", "E": "error"}


def test_analyze_links_is_the_function_behind_relations(tmp_path: Path) -> None:
    write(tmp_path, "a.csv", csv("customer_id;x", *[f"C-{n};{n}" for n in range(8)]))
    write(tmp_path, "b.csv", csv("customer_id;y", *[f"C-{n};{n}" for n in range(8)]))
    # a profile read back from JSON
    p = DataProfile.model_validate_json(profile(tmp_path).to_json())

    links = analyze_links(p)
    rebuilt = with_relations(p)

    assert links.relations == relations(p) and links.relations
    assert rebuilt.relations == links.relations


def test_two_columns_that_are_keys_of_nothing_never_reach_the_threshold(tmp_path: Path) -> None:
    values = [f"state {n % 6}" for n in range(12)]
    write(tmp_path, "a.csv", csv("status;x", *[f"{v};{n}" for n, v in enumerate(values)]))
    write(tmp_path, "b.csv", csv("status;y", *[f"{v};{n}" for n, v in enumerate(values)]))

    (found,) = profile(tmp_path).relations

    assert found.overlap == 1.0 and found.confidence == 0.7  # capped: neither side is a key


@pytest.mark.parametrize(
    ("table", "column"),
    [("categories", "category_id"), ("customers", "customer_id"), ("suppliers", "supplier_ref")],
)
def test_a_singular_column_name_points_at_its_plural_table_by_name(
    tmp_path: Path, table: str, column: str
) -> None:
    write(tmp_path, f"{table}.csv", csv("id;name", "1;a", "2;b", "3;c"))
    write(tmp_path, "items.csv", csv(f"sku;{column}", "S1;1", "S2;2", "S3;3"))

    [found] = relations(profile(tmp_path), values=lambda ids: {})

    ends = {found.from_column.rsplit(".", 1)[1], found.to_column.rsplit(".", 1)[1]}
    assert ends == {column, "id"} and found.kind == "name_match"
