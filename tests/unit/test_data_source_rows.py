"""Reading the real rows and text of a profiled file again, agreeing with the profile on which rows are data, with read-ahead."""

from __future__ import annotations

import datetime
import sqlite3
import threading
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import IO, Any

import pytest

from _helpers.data_fixtures import ORDERS_CSV, write_text
from awa.data import DataError, RowLimitError, _profile, profile
from awa.data._model import ColumnProfile, DataProfile, FileProfile, TableProfile, TextInfo
from awa.data._profile import _read_ahead
from awa.data._relations import analyze_links
from awa.data._source_rows import document_text, table_rows
from awa.data._values import local_values


def columns(*names: str) -> list[ColumnProfile]:
    return [ColumnProfile(name=n, key=n.lower(), index=i) for i, n in enumerate(names)]


def csv_file(path: Path, table: TableProfile, **info: object) -> FileProfile:
    return FileProfile(
        id="f1",
        path=path.name,
        source=str(path),
        sha256="0" * 64,
        size=path.stat().st_size,
        kind="csv",
        text=TextInfo(encoding="cp1252", delimiter=";", quote='"', **info),  # type: ignore[arg-type]
        tables=[table],
    )


def test_the_rows_of_a_real_german_csv_come_back_as_written() -> None:
    p = profile(ORDERS_CSV)
    file = p.files[0]

    rows = list(table_rows(file, file.tables[0], max_rows=100))

    assert len(rows) == file.tables[0].rows == 8
    assert rows[0] == [
        "B-1001",
        "Kowalski Stahl Sp. z o.o.",
        "PL",
        "4.250,00",
        "11.11.2024",
        "offen",
    ]
    assert all(len(row) == 6 for row in rows)


def test_title_lines_blank_rows_ragged_rows_and_totals_are_handled(tmp_path: Path) -> None:
    path = write_text(
        tmp_path,
        "t.csv",
        "Report 2026\r\nid;amount;note\r\n1;5,00;a\r\n\r\n2;7,50\r\n3;1,00;c;extra\r\nSumme;13,50;\r\n",
        "cp1252",
    )
    table = TableProfile(
        id="f1.t1",
        header_row=2,
        data_start_row=3,
        totals_rows=[7],
        columns=columns("id", "amount", "note"),
    )

    rows = list(table_rows(csv_file(path, table), table, max_rows=10))

    assert rows == [["1", "5,00", "a"], ["2", "7,50", None], ["3", "1,00", "c"]]


def test_a_table_without_a_header_starts_at_the_first_row(tmp_path: Path) -> None:
    path = write_text(tmp_path, "t.csv", "1;a\r\n2;b\r\n", "cp1252")
    table = TableProfile(id="f1.t1", header_row=0, data_start_row=1, columns=columns("A", "B"))

    assert list(table_rows(csv_file(path, table), table, max_rows=10)) == [["1", "a"], ["2", "b"]]


def test_too_many_rows_stop_the_reading_with_a_typed_error(tmp_path: Path) -> None:
    path = write_text(tmp_path, "t.csv", "a\r\n" + "1\r\n" * 50, "cp1252")
    table = TableProfile(id="f1.t1", columns=columns("a"))
    rows = table_rows(csv_file(path, table), table, max_rows=10)

    with pytest.raises(RowLimitError, match="more than 10 rows"):
        list(rows)


def test_a_missing_source_is_a_data_error_that_does_not_show_the_path(tmp_path: Path) -> None:
    path = write_text(tmp_path, "t.csv", "a\r\n1\r\n", "cp1252")
    table = TableProfile(id="f1.t1", columns=columns("a"))
    file = csv_file(path, table)
    path.unlink()

    with pytest.raises(DataError) as error:
        table_rows(file, table, max_rows=10)

    assert str(tmp_path) not in str(error.value)
    file.source = None
    with pytest.raises(DataError, match="source file is not available"):
        table_rows(file, table, max_rows=10)


def test_a_workbook_is_read_by_sheet_range_and_header_row(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Umsatz"
    sheet.append(["Report", None, None])
    sheet.append([None, None, None])
    sheet.append(["Kunde", "Betrag", "Datum"])
    sheet.append(["A", 10.5, datetime.date(2026, 1, 31)])
    sheet.append([None, None, None])
    sheet.append(["B", 4, datetime.datetime(2026, 2, 1, 8, 30)])
    sheet.append(["Summe", 14.5, None])
    workbook.create_sheet("Ziele").append(["x"])
    path = tmp_path / "w.xlsx"
    workbook.save(path)
    table = TableProfile(
        id="f1.t1",
        sheet="Umsatz",
        header_row=3,
        data_start_row=4,
        range="A3:C7",
        totals_rows=[7],
        columns=columns("Kunde", "Betrag", "Datum"),
    )
    file = FileProfile(
        id="f1",
        path="w.xlsx",
        source=str(path),
        sha256="0" * 64,
        size=1,
        kind="xlsx",
        tables=[table],
    )

    rows = list(table_rows(file, table, max_rows=10))

    assert rows == [["A", 10.5, "2026-01-31"], ["B", 4, "2026-02-01T08:30:00"]]


def test_a_sqlite_table_is_read_only_and_names_are_quoted(tmp_path: Path) -> None:
    path = tmp_path / "w.db"
    connection = sqlite3.connect(path)
    connection.execute('CREATE TABLE "odd ""name" (id INTEGER, "my col" TEXT)')
    connection.executemany('INSERT INTO "odd ""name" VALUES (?, ?)', [(1, "a"), (2, "b"), (3, "c")])
    connection.commit()
    connection.close()
    table = TableProfile(id="f1.t1", name='odd "name', columns=columns("id", "my col"))
    file = FileProfile(
        id="f1",
        path="w.db",
        source=str(path),
        sha256="0" * 64,
        size=1,
        kind="sqlite",
        tables=[table],
    )

    assert list(table_rows(file, table, max_rows=10)) == [[1, "a"], [2, "b"], [3, "c"]]
    with pytest.raises(RowLimitError):
        list(table_rows(file, table, max_rows=2))
    hostile = table.model_copy(update={"name": 'x"; DROP TABLE "odd ""name"; --'})
    with pytest.raises(DataError):
        list(table_rows(file, hostile, max_rows=10))
    assert sqlite3.connect(path).execute('SELECT count(*) FROM "odd ""name"').fetchone() == (3,)


def test_text_of_json_and_plain_files_and_the_whole_flag(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = write_text(tmp_path, "a.json", '{"id": "INV-77"}')
    file = FileProfile(
        id="f1", path="a.json", source=str(path), sha256="0" * 64, size=1, kind="json"
    )

    assert document_text(file) == ('{"id": "INV-77"}', True)
    monkeypatch.setattr("awa.data._source_rows.MAX_TEXT_CHARS", 5)
    assert document_text(file) == ('{"id"', False)
    with pytest.raises(DataError, match="has no text"):
        document_text(file.model_copy(update={"kind": "csv"}))


def test_a_damaged_word_file_is_a_data_error_not_a_crash(tmp_path: Path) -> None:
    path = write_text(tmp_path, "a.docx", "this is not a zip")
    file = FileProfile(
        id="f1", path="a.docx", source=str(path), sha256="0" * 64, size=1, kind="docx"
    )

    with pytest.raises(DataError, match="could not be read again"):
        document_text(file)


def mail_file(path: Path) -> FileProfile:
    return FileProfile(
        id="f1", path=path.name, source=str(path), sha256="0" * 64, size=1, kind="mail"
    )


def test_mail_text_is_decoded_like_a_mail_reader_shows_it(tmp_path: Path) -> None:
    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"] = "Rechnung für Müller"  # an encoded header
    message["From"] = "ann@example.com"
    message.set_content("Bitte zahlen: Größe 4.250,00 EUR, Referenz RE-77.", cte="quoted-printable")
    message.add_alternative("<p>Referenz <b>RE-77</b></p>", subtype="html")
    path = tmp_path / "m.eml"
    path.write_bytes(message.as_bytes())

    text, whole = document_text(mail_file(path))

    assert whole is True
    assert "Subject: Rechnung für Müller" in text and "Größe 4.250,00 EUR" in text
    assert "<b>RE-77</b>" in text


def test_a_mail_with_an_attachment_that_cannot_be_read_is_not_whole(tmp_path: Path) -> None:
    from email.message import EmailMessage

    message = EmailMessage()
    message.set_content("see attachment")
    message.add_attachment(
        b"\x00\x01binary", maintype="application", subtype="zip", filename="a.zip"
    )
    path = tmp_path / "m.eml"
    path.write_bytes(message.as_bytes())

    text, whole = document_text(mail_file(path))

    assert "see attachment" in text and whole is False


def test_a_mailbox_is_read_message_by_message_and_cut_at_the_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    raw = "".join(
        f"From a@example.com Mon Jan  1 00:00:00 2026\nSubject: n{n}\n\nbody {n}\n\n"
        for n in range(3)
    )
    path = write_text(tmp_path, "box.mbox", raw)

    text, whole = document_text(mail_file(path))
    assert whole is True and "body 0" in text and "body 2" in text

    monkeypatch.setattr("awa.data._source_rows.MAX_MESSAGES", 2)
    text, whole = document_text(mail_file(path))
    assert whole is False and "body 1" in text and "body 2" not in text


def _write(folder: Path, count: int) -> list[Path]:
    paths = []
    for number in range(count):
        path = folder / f"file_{number:02d}.csv"
        path.write_text(f"id;amount\n{number};1,5\n{number + 1};2,5\n", encoding="utf-8")
        paths.append(path)
    return paths


@pytest.fixture
def opened(monkeypatch: pytest.MonkeyPatch) -> list[Path]:
    """The paths opened by a read-ahead thread, in the order they were opened."""

    calls: list[Path] = []
    original = Path.open

    def spy(self: Path, *args: Any, **kwargs: Any) -> IO[Any]:
        if threading.current_thread().name == "awa-data-read-ahead":
            calls.append(self)
        return original(self, *args, **kwargs)  # type: ignore[no-any-return]

    monkeypatch.setattr(Path, "open", spy)
    return calls


def test_every_file_is_opened_once_by_a_read_ahead_thread(
    tmp_path: Path, opened: list[Path]
) -> None:
    paths = _write(tmp_path, 20)
    threads = _read_ahead(paths)
    for thread in threads:
        thread.join(5)
    assert 2 <= len(threads) <= 8
    assert sorted(opened) == paths


def test_a_file_that_cannot_be_opened_does_not_stop_the_others(
    tmp_path: Path, opened: list[Path]
) -> None:
    paths = _write(tmp_path, 6)
    paths.insert(2, tmp_path / "gone.csv")  # vanished since the folder was listed
    paths.insert(4, tmp_path)  # a folder cannot be read as a file
    for thread in _read_ahead(paths):
        thread.join(5)
    assert sorted(opened) == sorted(paths)


def test_one_file_needs_no_thread(tmp_path: Path) -> None:
    assert _read_ahead(_write(tmp_path, 1)) == []
    assert _read_ahead([]) == []


def test_the_profile_is_the_same_without_read_ahead(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _write(tmp_path, 10)
    (tmp_path / "broken.csv").write_bytes(b"\x00\x01\x02\x03" * 50)
    (tmp_path / "empty.csv").write_bytes(b"")
    with_it = profile(tmp_path, created="2026-01-01T00:00:00Z").model_dump_json()
    monkeypatch.setattr(_profile, "_read_ahead", lambda paths: [])
    without = profile(tmp_path, created="2026-01-01T00:00:00Z").model_dump_json()
    assert with_it == without


MULTILINE_TOTALS = 'id,note,amount\n1,"line one\nline two",10\n2,x,20\n3,y,30\nTotal,,60\n'
MULTILINE_ABOVE_HEADER = 'Report "A\nB"\nid,amount\n1,10\n2,20\n3,30\nTotal,60\n'
BLANK_LINES = "id,amount\n\n1,10\n\n\n2,20\n3,30\nTotal,60\n"
NEWLINES_EVERYWHERE = 'id,note,amount\n1,"a\n\nb",10\n2,"c\nd",20\n3,"e\nf\ng",30\nTotal,,60\n'
ONE_COLUMN = 'name\n"alpha, one"\n"beta, two"\n"gamma, three"\n"delta, four"\n'


def _rows(p: DataProfile) -> tuple[FileProfile, TableProfile, list[list[Any]]]:
    file = p.files[0]
    table = file.tables[0]
    return file, table, list(table_rows(file, table, max_rows=1000))


def _assert_same_data(p: DataProfile, *, column: str, expect: list[str]) -> None:
    _, table, rows = _rows(p)
    position = [c.key for c in table.columns].index(column)
    assert [str(row[position]) for row in rows] == expect, "table_rows disagrees with the profile"
    assert len(rows) == table.rows
    column_id = f"{table.id}.{column}"
    values = local_values(p)([column_id])[column_id]
    assert values.values == frozenset(expect), "local_values disagrees with the profile"
    assert values.cells == len(expect)


@pytest.mark.parametrize(
    "text",
    [MULTILINE_TOTALS, MULTILINE_ABOVE_HEADER, BLANK_LINES, NEWLINES_EVERYWHERE],
    ids=["multiline-cell-before-totals", "multiline-title", "blank-lines", "newlines-everywhere"],
)
def test_csv_readers_skip_the_rows_the_profile_skips(tmp_path: Path, text: str) -> None:
    p = profile(write_text(tmp_path, "t.csv", text))

    assert p.files[0].tables[0].totals_rows, "the fixture must have a totals row"
    _assert_same_data(p, column="id", expect=["1", "2", "3"])


def test_csv_totals_row_is_not_read_as_data_after_a_multiline_cell(tmp_path: Path) -> None:
    """The case of the audit: the totals row (line 6) must not come back as data."""

    p = profile(write_text(tmp_path, "t.csv", MULTILINE_TOTALS))
    _, table, rows = _rows(p)

    assert table.totals_rows == [6]
    assert ["Total", "", "60"] not in rows
    assert len(rows) == 3


def test_a_one_column_file_is_not_split_at_commas_by_any_reader(tmp_path: Path) -> None:
    p = profile(write_text(tmp_path, "t.csv", ONE_COLUMN))
    file, table, rows = _rows(p)

    assert file.text is not None and file.text.delimiter is None
    assert [row[0] for row in rows][:2] == ["alpha, one", "beta, two"]
    column_id = f"{table.id}.name"
    assert "alpha, one" in local_values(p)([column_id])[column_id].values


def test_csv_with_a_byte_order_mark_and_crlf_is_read_alike(tmp_path: Path) -> None:
    text = "﻿id;amount\r\n1;10\r\n2;20\r\n3;30\r\nSumme;60\r\n"
    p = profile(write_text(tmp_path, "t.csv", text, "utf-8"))

    _assert_same_data(p, column="id", expect=["1", "2", "3"])


def _book(tmp_path: Path, build: Callable[[Any], None], name: str = "t.xlsx") -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    build(workbook.active)
    path = tmp_path / name
    workbook.save(path)
    return path


def test_excel_readers_skip_the_rows_the_profile_skips(tmp_path: Path) -> None:
    def build(sheet: Any) -> None:
        sheet.append(["Report 2026"])
        sheet.append([])
        sheet.append(["id", "amount"])
        for n in (1, 2, 3):
            sheet.append([f"A-{n}", n * 10])
        sheet.append(["Total", 60])

    p = profile(_book(tmp_path, build))

    assert p.files[0].tables[0].totals_rows
    _assert_same_data(p, column="id", expect=["A-1", "A-2", "A-3"])


def test_excel_readers_do_not_trust_a_wrong_declared_dimension(tmp_path: Path) -> None:
    def build(sheet: Any) -> None:
        sheet.append(["id", "amount"])
        for n in (1, 2, 3, 4):
            sheet.append([f"A-{n}", n])

    source = _book(tmp_path, build)
    damaged = tmp_path / "damaged.xlsx"
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(damaged, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b'<dimension ref="A1:B5"/>' in data
                data = data.replace(b'<dimension ref="A1:B5"/>', b'<dimension ref="A1:B2"/>')
            out.writestr(item, data)
    p = profile(damaged)

    _assert_same_data(p, column="id", expect=["A-1", "A-2", "A-3", "A-4"])


def test_excel_second_table_on_a_sheet_is_read_from_its_own_range(tmp_path: Path) -> None:
    def build(sheet: Any) -> None:
        sheet.append(["id", "amount"])
        sheet.append(["A-1", 1])
        sheet.append(["A-2", 2])
        sheet.append([])
        sheet.append([])
        sheet.append(["code", "price"])
        sheet.append(["Z-1", 1.5])
        sheet.append(["Z-2", 2.5])

    p = profile(_book(tmp_path, build))
    file = p.files[0]
    assert len(file.tables) == 2
    second = file.tables[1]
    rows = list(table_rows(file, second, max_rows=100))
    column_id = f"{second.id}.code"

    assert [r[0] for r in rows] == ["Z-1", "Z-2"]
    assert local_values(p)([column_id])[column_id].values == frozenset({"Z-1", "Z-2"})


def test_sqlite_reader_returns_every_row_the_profile_counts(tmp_path: Path) -> None:
    database = tmp_path / "t.db"
    with sqlite3.connect(database) as connection:
        connection.execute("CREATE TABLE t (id TEXT, amount INTEGER)")
        connection.executemany("INSERT INTO t VALUES (?, ?)", [(f"S-{n}", n) for n in range(5)])
    connection.close()
    _assert_same_data(profile(database), column="id", expect=[f"S-{n}" for n in range(5)])


def test_a_sqlite_file_in_a_folder_with_uri_characters_is_read(tmp_path: Path) -> None:
    """`#` and `%` (and `?`, where the system allows it) in a path must not break the read-only URI (the relations lost evidence)."""

    folder = tmp_path / "a#b%20c"
    folder.mkdir()
    database = folder / "f1.db"
    connection = sqlite3.connect(database)
    connection.execute("CREATE TABLE t1 (id TEXT)")
    connection.executemany("INSERT INTO t1 VALUES (?)", [(f"K-{n}",) for n in range(5)])
    connection.commit()
    connection.close()
    p = profile(database)
    column_id = f"{p.files[0].tables[0].id}.id"

    assert local_values(p)([column_id])[column_id].values == frozenset(f"K-{n}" for n in range(5))


def _workbook_with_ids(tmp_path: Path, name: str) -> Path:
    def build(sheet: Any) -> None:
        sheet.append(["customer_id", "name"])
        for n in range(1, 9):
            sheet.append([f"C-{n}", f"name {n}"])

    return _book(tmp_path, build, name)


def test_columns_that_cannot_be_read_again_are_reported(tmp_path: Path) -> None:
    _workbook_with_ids(tmp_path, "orders.xlsx")
    _workbook_with_ids(tmp_path, "customers.xlsx")
    p = profile(tmp_path)
    first = p.files[0].tables[0]
    source = Path(p.files[0].source or "")
    source.write_bytes(b"\x00" * source.stat().st_size)  # same size, no longer a workbook
    failures: list[tuple[str, str]] = []

    read = local_values(p, on_error=lambda table_id, error: failures.append((table_id, error)))
    found = read([f"{first.id}.customer_id"])

    assert found == {}
    assert failures and failures[0][0] == first.id and "BadZipFile" in failures[0][1]
    links = analyze_links(p)
    assert any(flag.code == "DF-W-VALUES-UNREAD" for flag in links.flags)


def test_a_source_that_is_gone_is_not_an_error(tmp_path: Path) -> None:
    _workbook_with_ids(tmp_path, "orders.xlsx")
    _workbook_with_ids(tmp_path, "customers.xlsx")
    p = profile(tmp_path)
    for file in p.files:
        file.source = None
    failures: list[tuple[str, str]] = []

    local_values(p, on_error=lambda table_id, error: failures.append((table_id, error)))(
        [f"{p.files[0].tables[0].id}.customer_id"]
    )

    assert failures == []
    assert not any(flag.code == "DF-W-VALUES-UNREAD" for flag in analyze_links(p).flags)
