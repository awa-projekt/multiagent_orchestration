"""Shared helpers of the automation use case tests: project copies, a served worker,
scripted models, and fake outside systems. Imported by the `test_usecase_*.py` modules."""

from __future__ import annotations

import asyncio
import json
import re
import shutil
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from _helpers.conversation import results, system_text
from _helpers.framing import tool_content
from _helpers.polling import poll
from awa import Blueprint
from awa.automation import SqliteRunStore
from awa.blueprints.json_schema import model_from_json_schema
from awa.cli import build_parser
from awa.models import ModelRegistry
from awa.server import build_blueprint_app
from awa.testing import FunctionChatModel, ModelCall, structured_output

USECASES = Path(__file__).resolve().parents[2] / "examples" / "usecases"
AUTOMATION_USECASES = [
    "compliance-news-watch",
    "data-quality-reconciliation",
    "email-inbox-triage",
    "weekly-management-report",
    "inventory-replenishment",
    "incident-postmortem",
]
TEMPLATE_USECASES = ("email-inbox-triage", "expense-report-audit")
"""Use cases held to the template standard: 8 to 12 cases, two of them at a boundary."""
# The data-based evaluation use cases: expected answers computed from their data by
# `make_cases.py` (checked in `test_usecase_data_cases.py`).
DATA_USECASES = ["sales-operations"]


# ---------------------------------------------------------------------- project and server


def usecase(name: str) -> Any:
    """A fixture: a copy of the use case folder `name`, as the working directory."""

    @pytest.fixture
    def project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
        folder = tmp_path / name
        shutil.copytree(
            USECASES / name, folder, ignore=shutil.ignore_patterns("__pycache__", "out", ".awa")
        )
        monkeypatch.chdir(folder)
        monkeypatch.setattr(sys, "path", list(sys.path))
        monkeypatch.delitem(sys.modules, "tools", raising=False)
        monkeypatch.delenv("AWA_POSTGRES_DSN", raising=False)
        return folder

    return project


def cli(*argv: str) -> int:
    args = build_parser(list(argv)).parse_args(argv)
    code: int = args.handler(args)
    return code


def serve(project: Path) -> TestClient:
    """What `awa serve blueprint.yaml --tools tools.py` runs, in this process."""

    return TestClient(
        build_blueprint_app(
            "blueprint.yaml", tools=["tools.py"], store_path=project / ".awa" / "awa.db"
        )
    )


def runs_of(client: TestClient) -> list[dict[str, Any]]:
    runs: list[dict[str, Any]] = client.get("/v1/process-runs").json()["runs"]
    return runs


def _content(path: Path) -> object:
    """What a data file says. SQLite databases, workbooks and Word files record the version of
    the library that wrote them, so their bytes differ between machines; their content does not.
    Text data compare by their lines, so a checkout with CRLF line endings equals one with LF."""

    if path.suffix == ".db":
        import sqlite3
        from contextlib import closing

        with closing(sqlite3.connect(path)) as database:
            return list(database.iterdump())
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True)
        try:
            return [
                (sheet.title, [[cell.value for cell in row] for row in sheet.iter_rows()])
                for sheet in workbook.worksheets
            ]
        finally:
            workbook.close()
    if path.suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        return [p.text for p in document.paragraphs] + [
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        ]
    data = path.read_bytes()
    if path.suffix in (".csv", ".json", ".md", ".txt", ".log", ".yaml", ".yml"):
        return data.replace(b"\r\n", b"\n")
    return data


def data_differences(
    generated: Path, other: Path, *, ignore: tuple[str, ...], same_bytes: bool
) -> list[str]:
    """The files that differ between two data folders (empty folders do not count, as git keeps
    none). `same_bytes=False` compares what the files say, for data committed on another machine."""

    def files(root: Path) -> dict[str, Path]:
        return {
            found.relative_to(root).as_posix(): found
            for found in sorted(root.rglob("*"))
            if found.is_file() and not set(found.relative_to(root).parts) & set(ignore)
        }

    left, right = files(generated), files(other)
    problems = [f"only in {generated.name}: {name}" for name in left.keys() - right.keys()]
    problems += [f"only in {other.name}: {name}" for name in right.keys() - left.keys()]
    for name in sorted(left.keys() & right.keys()):
        if same_bytes:
            equal = left[name].read_bytes() == right[name].read_bytes()
        else:
            equal = _content(left[name]) == _content(right[name])
        if not equal:
            problems.append(f"differs: {name}")
    return problems


def settle(
    client: TestClient, status: str, *, count: int = 1, trigger: str | None = None
) -> list[dict[str, Any]]:
    """Wait until `count` runs (of `trigger`) have `status`; return them."""

    return poll(
        lambda: [
            r
            for r in runs_of(client)
            if r["status"] == status and (trigger is None or r["trigger"] == trigger)
        ],
        lambda found: len(found) >= count,
        timeout=30,
        what=f"{count} run(s) reaching {status!r}",
    )


def decide(run_id: str, *, approve: bool = True) -> None:
    flag = "--approve" if approve else "--reject"
    assert cli("runs", "decide", run_id, flag, "--actor", "lead") == 0


def seed_schedule(project: Path, trigger: str, name: str) -> None:
    """Pretend the worker last ticked ten days ago, so the next start catches up once."""

    async def seed() -> None:
        store = SqliteRunStore(project / ".awa" / "awa.db")
        past = (datetime.now(UTC) - timedelta(days=10)).isoformat()
        assert await store.compare_and_set_state(f"schedule:{name}:{trigger}", None, past)
        await store.close()

    asyncio.run(seed())


# ---------------------------------------------------------------------- scripted models


class Script:
    """One scripted model for every agent of a blueprint; one method per agent name."""

    def __init__(self, blueprint: str = "blueprint.yaml") -> None:
        self.blueprint = Blueprint.from_yaml(blueprint)
        self.markers = {
            spec.name: spec.instructions.strip().splitlines()[0] for spec in self.blueprint.agents
        }

    def install(self, monkeypatch: pytest.MonkeyPatch) -> None:
        registry = ModelRegistry()
        registry.register("default", FunctionChatModel(self))
        monkeypatch.setattr(ModelRegistry, "_default", registry)

    def __call__(self, call: ModelCall) -> Any:
        prompt = system_text(call)
        found = [name for name, marker in self.markers.items() if marker in prompt]
        assert found, f"cannot tell the agent from the prompt: {prompt[:200]!r}"
        return getattr(self, max(found, key=lambda name: len(self.markers[name])))(call)

    def finish(self, agent: str, data: dict[str, Any]) -> Any:
        schema = self.blueprint.agent(agent).output_schema
        assert schema is not None
        name = model_from_json_schema(schema, name=f"{agent}_output").__name__
        return structured_output(name, data)


def delegated(call: ModelCall) -> dict[str, list[str]]:
    """Results of every `delegate_task` call, by member name."""

    found: dict[str, list[str]] = {}
    pattern = r'<result agent="([\w-]+)"[^>]*>\s*(.*?)\s*</result>'
    for content in results(call, "delegate_task"):
        for name, text in re.findall(pattern, content, re.S):
            found.setdefault(name, []).append(text)
    return found


def first_json(text: str) -> Any:
    """The first JSON value embedded in `text` (a framed tool output is read inside its frame)."""

    text = tool_content(text)
    start = min(i for i in (text.find("{"), text.find("[")) if i >= 0)
    return json.JSONDecoder().raw_decode(text[start:])[0]


def lines(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def run_input(project: Path, run_id: str) -> str:
    """The input a run was started with (read from the local durable store)."""

    async def read() -> Any:
        store = SqliteRunStore(project / ".awa" / "awa.db")
        try:
            return await store.get(run_id)
        finally:
            await store.close()

    return str(asyncio.run(read()).input)


class FakeImap:
    """Enough of `imaplib.IMAP4` for the mailbox trigger: unseen messages, flags, fetch."""

    capabilities = ("IMAP4REV1",)

    def __init__(self, messages: dict[str, bytes]) -> None:
        self.messages = messages
        self.seen: set[str] = set()

    def login(self, user: str, password: str) -> tuple[str, list[bytes]]:
        return "OK", [b""]

    def select(self, folder: str) -> tuple[str, list[bytes]]:
        return "OK", [b"1"]

    def noop(self) -> tuple[str, list[bytes]]:
        return "OK", [b""]

    def uid(self, command: str, *args: Any) -> tuple[str, list[Any]]:
        if command == "SEARCH":
            return "OK", [" ".join(u for u in self.messages if u not in self.seen).encode()]
        if command == "FETCH":
            uid = str(args[0])
            return "OK", [(f"{uid} (BODY[] {{1}}".encode(), self.messages[uid]), b")"]
        if command == "STORE":
            self.seen.add(str(args[0]))
        return "OK", [b""]

    def expunge(self) -> tuple[str, list[bytes]]:
        return "OK", [b""]

    def logout(self) -> tuple[str, list[bytes]]:
        return "BYE", [b""]


def mail(sender: str, subject: str, text: str, message_id: str) -> bytes:
    """A plain-text e-mail as the IMAP server would return it."""

    from email.message import EmailMessage

    message = EmailMessage()
    message["From"], message["To"] = sender, "inbox@example.com"
    message["Subject"], message["Message-ID"] = subject, message_id
    message["Date"] = "Mon, 05 Oct 2026 08:15:00 +0200"
    message.set_content(text)
    return message.as_bytes()
