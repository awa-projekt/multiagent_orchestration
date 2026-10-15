"""Tools of the order-to-cash-exceptions use case: the three-way comparison and the dunning letter.

The agent explores the data with the `workspace` connector (SQL on the order database, Excel
invoices, CSV payments, emails); this file holds the real logic: `classify_order` compares an
order with its invoices and payments from the three sources and classifies the mismatch with
a recommended action, and `send_dunning_letter` queues a letter only for an invoice that is
really open, overdue and free of billing errors, and only with the right amount in the text.
In production they read your ERP and send through your dunning process (see the README).
"""

from __future__ import annotations

import csv
import json
import re
import sqlite3
from contextlib import closing
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from awa.runtime import case_dir

HERE = Path(__file__).resolve().parent


def _data_dir() -> Path:
    """The sample data: the eval case's copy inside `awa eval`, else the folder here."""

    return case_dir(HERE) / "data"


def _outbox_dir() -> Path:
    return _data_dir() / "outbox"


AS_OF: date | None = date(2026, 10, 4)
"""The reporting date; the sample data is dated around it. Set `None` to use today."""
MINOR_EUR = 25.0
"""Open differences up to this amount are written off, not chased."""
BILLING_ISSUES = {"invoiced_more_than_shipped", "duplicate_invoice", "invoiced_not_shipped"}
"""While one of these is open, no dunning letter goes out: the invoice has to be corrected first."""


def _query(sql: str, *args: Any) -> list[tuple[Any, ...]]:
    with closing(
        sqlite3.connect(f"file:{(_data_dir() / 'o2c.db').as_posix()}?mode=ro", uri=True)
    ) as db:
        return db.execute(sql, args).fetchall()


def _invoices() -> list[dict[str, Any]]:
    book = load_workbook(_data_dir() / "invoices.xlsx", read_only=True, data_only=True)
    sheet = book["Invoices"]
    head, *body = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    return [dict(zip(head, row, strict=True)) for row in body if row[0]]


def _payments() -> list[dict[str, str]]:
    with (_data_dir() / "payments.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _paid_by_invoice(
    invoices: list[dict[str, Any]],
) -> tuple[dict[str, float], list[dict[str, str]]]:
    """Payments matched to invoices by the invoice number in the reference (INV-3002, RE 3002,
    inv 3002), and the payments that match none."""

    numbers = {str(i["invoice_no"]).split("-")[-1]: str(i["invoice_no"]) for i in invoices}
    paid: dict[str, float] = {}
    unmatched = []
    for payment in _payments():
        hit = next((n for n in re.findall(r"\d{4,}", payment["reference"]) if n in numbers), None)
        if hit is None:
            unmatched.append(payment)
        else:
            paid[numbers[hit]] = paid.get(numbers[hit], 0.0) + float(payment["amount_eur"])
    return paid, unmatched


def _classify(order_id: str) -> dict[str, Any] | None:
    order = _query(
        "SELECT o.order_id, o.customer_id, c.name, c.email, o.shipped_eur FROM orders o "
        "JOIN customers c USING (customer_id) WHERE o.order_id = ?",
        order_id.strip().upper(),
    )
    if not order:
        return None
    order_id, customer_id, customer, email, shipped = order[0]
    all_invoices = _invoices()
    paid_by, _ = _paid_by_invoice(all_invoices)
    mine = [i for i in all_invoices if i["order_id"] == order_id]
    unique, seen = [], set()
    for invoice in mine:
        if invoice["amount_eur"] in seen:
            continue
        seen.add(invoice["amount_eur"])
        unique.append(invoice)
    billed = sum(i["amount_eur"] for i in unique)
    paid = sum(paid_by.get(i["invoice_no"], 0.0) for i in unique)
    open_eur = round(billed - paid, 2)
    issues: list[str] = []
    if not mine and shipped > 0:
        issues.append("shipped_not_invoiced")
    if mine and shipped == 0:
        issues.append("invoiced_not_shipped")
    elif mine and billed != shipped:
        issues.append(
            "invoiced_more_than_shipped" if billed > shipped else "invoiced_less_than_shipped"
        )
    if len(mine) > len(unique):
        issues.append("duplicate_invoice")
    if open_eur < 0:
        issues.append("overpaid")
    elif 0 < open_eur <= MINOR_EUR and paid > 0:
        issues.append("short_payment_minor")
    elif open_eur > 0 and paid > 0:
        issues.append("underpaid")
    due = min((str(i["due_date"]) for i in unique), default=None)
    days = (AS_OF or date.today()) - date.fromisoformat(due) if due else None
    collect = open_eur > MINOR_EUR or (open_eur > 0 and paid == 0)
    overdue = max(days.days, 0) if days is not None and collect else 0
    if collect and overdue:
        issues.append("unpaid_overdue" if paid == 0 else "overdue")
    elif collect:
        issues.append("unpaid_not_due")
    level = (
        0
        if not overdue or BILLING_ISSUES & set(issues)
        else 1
        if overdue <= 30
        else 2
        if overdue <= 60
        else 3
    )
    actions = {
        "shipped_not_invoiced": "issue the invoice",
        "invoiced_not_shipped": "cancel the invoice until the goods ship",
        "invoiced_more_than_shipped": f"credit the difference of {billed - shipped:.2f} EUR before any dunning",
        "duplicate_invoice": "cancel the duplicate invoice with a credit note",
        "overpaid": f"refund or credit {-open_eur:.2f} EUR",
        "short_payment_minor": f"write off {open_eur:.2f} EUR (below {MINOR_EUR:.0f} EUR)",
    }
    steps = [actions[i] for i in issues if i in actions]
    if level:
        steps.append(f"send dunning level {level}")
    return {
        "order_id": order_id,
        "customer_id": customer_id,
        "customer": customer,
        "contact": email,
        "shipped_eur": shipped,
        "invoices": [
            {
                "invoice_no": i["invoice_no"],
                "amount_eur": i["amount_eur"],
                "due_date": str(i["due_date"]),
            }
            for i in mine
        ],
        "billed_eur": billed,
        "paid_eur": round(paid, 2),
        "open_eur": open_eur,
        "days_overdue": overdue,
        "issues": issues or ["ok"],
        "dunning_level": level,
        "recommended_action": "; ".join(steps) or "nothing to do",
    }


def classify_order(order_id: str) -> str:
    """Compare an order with its invoices (Excel) and payments (CSV) and classify the result.

    Returns the amounts (shipped, billed, paid, open), the days overdue, the issues
    (ok, shipped_not_invoiced, invoiced_not_shipped, invoiced_more_than_shipped,
    invoiced_less_than_shipped, duplicate_invoice, overpaid, short_payment_minor, underpaid,
    unpaid_overdue, overdue, unpaid_not_due), the dunning level that is due (0 for none; 1 to 30
    days late, 2 to 60, 3 beyond; never while a billing error is open) and the recommended action.

    In production: read order, invoice and payment data from your ERP.
    """

    found = _classify(order_id)
    return json.dumps(found) if found else f"Unknown order {order_id!r}."


def find_unmatched_payments(customer_id: str) -> str:
    """List the payments of a customer whose reference names no invoice (cannot be allocated).

    In production: read the bank statement lines your cash application could not match.
    """

    _, unmatched = _paid_by_invoice(_invoices())
    mine = [p for p in unmatched if p["customer_id"] == customer_id.strip().upper()]
    return json.dumps(mine)


def send_dunning_letter(invoice_no: str, level: int, subject: str, body: str) -> str:
    """Queue a dunning letter to the customer of an invoice. Refused unless the invoice is open,
    overdue at that level and free of billing errors, and unless the text states the open
    amount. A person approves every letter before it goes out.

    In production: hand the letter to your dunning process or mail gateway.
    """

    invoice_no = invoice_no.strip().upper()
    order = next((i["order_id"] for i in _invoices() if i["invoice_no"] == invoice_no), None)
    state = _classify(order) if order else None
    if state is None:
        return f"Refused: unknown invoice {invoice_no!r}."
    if state["dunning_level"] == 0 or int(level) != state["dunning_level"]:
        return (
            f"Refused: {invoice_no} is not due for a level {level} letter "
            f"(issues: {', '.join(state['issues'])}; due level {state['dunning_level']})."
        )
    amount = state["open_eur"]
    spellings = (
        f"{amount:.2f}",
        f"{amount:,.2f}",
        f"{amount:,.2f}".translate(str.maketrans(",.", ".,")),
    )
    if not any(text in body for text in spellings):
        return f"Refused: the letter must state the open amount {amount:.2f} EUR."
    letter = _outbox_dir() / "dunning" / f"{invoice_no}-L{level}.txt"
    if letter.is_file():
        return f"Already queued: {letter.name}."
    letter.parent.mkdir(parents=True, exist_ok=True)
    subject = " ".join(subject.split())  # one header line, whatever the caller passes
    letter.write_text(f"To: {state['contact']}\nSubject: {subject}\n\n{body}\n", encoding="utf-8")
    return f"Queued level {level} letter for {invoice_no} to {state['contact']}."


__all__ = ["classify_order", "find_unmatched_payments", "send_dunning_letter"]
